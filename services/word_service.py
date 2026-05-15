from __future__ import annotations

from datetime import date
from io import BytesIO

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_font(run, name: str, size_pt: int, bold: bool = False, color: str | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.start_type = WD_SECTION_START.NEW_PAGE

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)


def build_report_docx(
    dataframe: pd.DataFrame,
    week_start: date,
    week_end: date,
    department_name: str | None,
    cell_name: str | None,
) -> bytes:
    document = Document()
    _configure_document(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run("주간업무보고")
    _set_font(title_run, "Calibri", 20, bold=True, color="143550")

    meta = document.add_paragraph()
    meta_run = meta.add_run(
        f"보고기간 {week_start.isoformat()} ~ {week_end.isoformat()}  |  "
        f"부서 {department_name or '전체'}  |  셀 {cell_name or '전체'}"
    )
    _set_font(meta_run, "Calibri", 10, color="4B6173")

    summary = document.add_paragraph()
    summary_run = summary.add_run(f"총 {len(dataframe)}건의 업무 항목이 조회되었습니다.")
    _set_font(summary_run, "Calibri", 11, color="1F4D78")

    grouped = dataframe.groupby(["부서", "셀", "직번", "이름"], sort=False)
    for department, cell, emp_no, name in grouped.groups.keys():
        heading = document.add_paragraph()
        heading_run = heading.add_run(f"{name} ({emp_no})")
        _set_font(heading_run, "Calibri", 14, bold=True, color="1F4D78")

        detail = document.add_paragraph()
        detail_run = detail.add_run(f"부서 {department}  |  셀 {cell}")
        _set_font(detail_run, "Calibri", 10, color="4B6173")

        rows = grouped.get_group((department, cell, emp_no, name)).reset_index(drop=True)
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        table.autofit = False

        widths = [Inches(2.2), Inches(0.9), Inches(2.5), Inches(0.9)]
        headers = ["SR 제목", "진행률", "SR 개발내용", "작성일시"]

        header_cells = table.rows[0].cells
        for idx, (cell_obj, header, width) in enumerate(zip(header_cells, headers, widths)):
            cell_obj.width = width
            paragraph = cell_obj.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(header)
            _set_font(run, "Calibri", 10, bold=True, color="FFFFFF")
            cell_obj.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_shading(cell_obj, "1F4D78")

        for row in rows.itertuples(index=False):
            table_row = table.add_row().cells
            values = [
                str(row[6] or ""),
                str(row[7] or ""),
                str(row[8] or ""),
                str(row[9] or ""),
            ]
            for idx, (cell_obj, value, width) in enumerate(zip(table_row, values, widths)):
                cell_obj.width = width
                paragraph = cell_obj.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in {1, 3} else WD_ALIGN_PARAGRAPH.LEFT
                run = paragraph.add_run(value)
                _set_font(run, "Calibri", 10)
                cell_obj.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        document.add_paragraph()

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
